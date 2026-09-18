from celery import shared_task
from app.core.database import async_session_maker
from app.models import Export, Job, JobContact, ContactType, JobVerification, User
from sqlalchemy import select
from sqlalchemy.orm import selectinload
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import uuid
from datetime import datetime, timedelta
import structlog

logger = structlog.get_logger()


EXPORT_COLUMNS = [
    ("Job ID", "id"),
    ("Job Title", "title"),
    ("Company", "company_name"),
    ("Source", "source_name"),
    ("Source URL", "source_url"),
    ("Location", "location"),
    ("Employment Type", "employment_type"),
    ("Salary Min", "salary_min"),
    ("Salary Max", "salary_max"),
    ("Currency", "currency"),
    ("Posted Date", "posted_at"),
    ("Recruiter", "recruiter_name"),
    ("Email", "email"),
    ("Phone", "phone"),
    ("Employer Verification", "employer_verification"),
    ("Uploader Verification", "uploader_verification"),
    ("Contact Verification", "contact_verification"),
    ("Risk Level", "risk_level"),
    ("Risk Score", "risk_score"),
    ("Final Confidence", "final_confidence"),
    ("Verification Status", "status"),
    ("Verification Date", "verified_at"),
]


@shared_task(bind=True, max_retries=2, default_retry_delay=120)
def generate_export_task(self, export_id: int):
    import asyncio
    return asyncio.run(_generate_export_async(export_id))


async def _generate_export_async(export_id: int):
    async with async_session_maker() as db:
        stmt = select(Export).where(Export.id == export_id)
        result = await db.execute(stmt)
        export = result.scalar_one_or_none()

        if not export:
            logger.error("export_not_found", export_id=export_id)
            return {"status": "error", "message": "Export not found"}

        export.status = "processing"
        await db.commit()

        try:
            jobs = await _fetch_jobs_for_export(db, export.filters, export.user_id)

            filename = f"JobShield_{datetime.utcnow().strftime('%Y-%m-%d')}_{uuid.uuid4().hex[:8]}.xlsx"
            filepath = os.path.join("/tmp/exports", filename)
            os.makedirs(os.path.dirname(filepath), exist_ok=True)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Jobs"

            _write_header(ws)
            _write_data(ws, jobs)
            _style_worksheet(ws, len(jobs) + 1)

            wb.save(filepath)

            export.file_path = filepath
            export.row_count = len(jobs)
            export.status = "completed"
            export.completed_at = datetime.utcnow()
            export.expires_at = datetime.utcnow() + timedelta(hours=24)
            export.download_url = f"/api/v1/exports/{export.id}/download"
            await db.commit()

            logger.info("export_completed", export_id=export_id, rows=len(jobs), file=filename)
            return {"status": "success", "export_id": export_id, "rows": len(jobs), "file": filename}

        except Exception as e:
            export.status = "failed"
            export.error_message = str(e)
            await db.commit()
            logger.error("export_failed", export_id=export_id, error=str(e))
            raise


async def _fetch_jobs_for_export(db, filters: dict, user_id: int):
    stmt = (
        select(Job)
        .where(Job.user_id == user_id)
        .options(
            selectinload(Job.company),
            selectinload(Job.recruiter),
            selectinload(Job.contacts),
            selectinload(Job.verification),
            selectinload(Job.source),
        )
    )

    if filters.get("status"):
        stmt = stmt.where(Job.status == filters["status"])
    if filters.get("risk_level"):
        stmt = stmt.where(Job.risk_level == filters["risk_level"])
    if filters.get("source_id"):
        stmt = stmt.where(Job.source_id == filters["source_id"])
    if filters.get("has_email"):
        stmt = stmt.where(Job.contacts.any(JobContact.type == ContactType.EMAIL))
    if filters.get("has_phone"):
        stmt = stmt.where(Job.contacts.any(JobContact.type == ContactType.PHONE))
    if filters.get("company_id"):
        stmt = stmt.where(Job.company_id == filters["company_id"])
    if filters.get("date_from"):
        stmt = stmt.where(Job.posted_at >= filters["date_from"])
    if filters.get("date_to"):
        stmt = stmt.where(Job.posted_at <= filters["date_to"])

    stmt = stmt.order_by(Job.created_at.desc())
    result = await db.execute(stmt)
    return result.scalars().all()


def _write_header(ws):
    for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B1E22", end_color="1B1E22", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _write_data(ws, jobs):
    for row_idx, job in enumerate(jobs, 2):
        email = ""
        phone = ""
        for contact in job.contacts:
            if contact.type == ContactType.EMAIL and contact.is_valid:
                email = contact.normalized_value
                break
        for contact in job.contacts:
            if contact.type == ContactType.PHONE and contact.is_valid:
                phone = contact.normalized_value
                break

        verification = job.verification
        employer_ver = "Verified" if verification and verification.employer_score > 0.7 else "Not Verified"
        uploader_ver = "Verified" if verification and verification.uploader_score > 0.7 else "Not Verified"
        contact_ver = "Verified" if verification and verification.contact_score > 0.7 else "Not Verified"

        row_data = [
            job.id,
            job.title,
            job.company.name if job.company else "",
            job.source.name if job.source else "",
            job.source_url,
            job.location or "",
            job.employment_type or "",
            job.salary_min or "",
            job.salary_max or "",
            job.currency,
            job.posted_at.strftime("%Y-%m-%d") if job.posted_at else "",
            job.recruiter.name if job.recruiter else "",
            email,
            phone,
            employer_ver,
            uploader_ver,
            contact_ver,
            job.risk_level.value if job.risk_level else "",
            f"{verification.risk_score:.2f}" if verification else "",
            f"{verification.final_confidence:.2f}" if verification else "",
            job.status.value if job.status else "",
            job.verified_at.strftime("%Y-%m-%d %H:%M") if job.verified_at else "",
        ]

        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def _style_worksheet(ws, max_row):
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=len(EXPORT_COLUMNS)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    column_widths = [10, 35, 25, 15, 40, 20, 18, 12, 12, 8, 12, 25, 35, 18, 18, 18, 18, 12, 10, 12, 18, 18]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


@shared_task
def cleanup_expired_exports():
    import asyncio
    return asyncio.run(_cleanup_expired_exports_async())


async def _cleanup_expired_exports_async():
    async with async_session_maker() as db:
        stmt = select(Export).where(Export.expires_at < datetime.utcnow(), Export.status == "completed")
        result = await db.execute(stmt)
        exports = result.scalars().all()

        count = 0
        for export in exports:
            if export.file_path and os.path.exists(export.file_path):
                os.remove(export.file_path)
            export.status = "expired"
            count += 1

        await db.commit()
        return {"cleaned": count}